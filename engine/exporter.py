"""文书导出：把系统里的事实落成可打印的 Word / PPT / eml 文件。

设计原则
- 内容不重算：letters.py 已经产出「事实 → 依据 → 要求」的合规文本，
  这里只做排版（markdown → docx），保证文书与对话里的回答一字不差。
- 样式统一：页眉校名、标题黑体、落款日期，出来的文件可以直接打印盖章。
- 邮件双模式：配置了 SMTP（AG_SMTP_*）就真实发送；没配置就产出 .eml，
  收件人已填好，本地双击即进邮件客户端——演示环境同样能走完闭环。
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SCHOOL_NAME = "传智理工学院"
INK = RGBColor(0x3B, 0x2E, 0x24)  # 主文字色（暖棕，与前端 --text 一致）
BRAND = RGBColor(0xC2, 0x41, 0x0C)  # 品牌色（暖赭橙，与前端品牌同族）
MUTED = RGBColor(0x8E, 0x7A, 0x66)  # 次要文字色（暖灰棕）

# ---------------------------------------------------------------------------
# 字体方案：正文宋体 + 标题黑体。两者都是 Windows 系统自带字体，
# Word 与 WPS 在任何机器上打开都直接命中、渲染一致（等线是 Office 2016+
# 附带字体，WPS 常缺失会静默回退，故弃用）。PPT 用微软雅黑同理。
# ---------------------------------------------------------------------------
FONT_BODY = "宋体"
FONT_HEAD = "黑体"
FONT_PPT = "微软雅黑"


# ---------------------------------------------------------------------------
# markdown → docx：letters.py 的输出是受限 markdown（粗体/三级标题/列表/表格），
# 逐行解析足以覆盖，不需要通用 markdown 库。
# ---------------------------------------------------------------------------

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_runs(par, text: str, base_size=10.5):
    """把 **bold** 拆成多个 run 写进段落，每个 run 都显式钉死字体。"""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            _set_font(par.add_run(text[pos:m.start()]), size=base_size)
        _set_font(par.add_run(m.group(1)), size=base_size, bold=True)
        pos = m.end()
    if pos < len(text):
        _set_font(par.add_run(text[pos:]), size=base_size)


def _set_font(run, name=FONT_BODY, size=10.5, color=INK, bold=False):
    """统一字体入口：ascii/hAnsi/eastAsia/cs 四处一起设，Word 与 WPS 渲染一致。"""
    run.font.name = name  # 设 w:ascii + w:hAnsi（西文）
    rFonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)  # 中文字形（Word/WPS 都读这里）
    rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def _new_doc(header_text: str) -> Document:
    doc = Document()
    # 全局兜底：Normal 样式钉死宋体——表格单元格、空段落等未显式设字体处全部继承
    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(10.5)
    nrFonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        nrFonts.set(qn(attr), FONT_BODY)
    # 页面边距：正式文书的留白
    for sec in doc.sections:
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.6)
        sec.right_margin = Cm(2.6)
    # 页眉：校名 + 文书标识
    hp = doc.sections[0].header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(f"{SCHOOL_NAME} · {header_text}")
    _set_font(r, size=8, color=MUTED)
    return doc


def _title(doc: Document, text: str, sub: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    _set_font(r, name=FONT_HEAD, size=16, bold=True)
    if sub:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(sub)
        _set_font(r2, size=10, color=MUTED)
    # 标题下加一条细横线，文书感
    p3 = doc.add_paragraph()
    pr = p3._p.get_or_add_pPr()
    bdr = pr.makeelement(qn("w:pBdr"), {})
    bottom = pr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "6",
                                             qn("w:space"): "1", qn("w:color"): "C2410C"})
    bdr.append(bottom)
    pr.append(bdr)


def md_to_docx(doc: Document, body: str):
    """letters.py 的受限 markdown → Word 段落。"""
    lines = body.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        # 三级标题
        if line.startswith("### "):
            p = doc.add_paragraph()
            r = p.add_run(line[4:])
            _set_font(r, name=FONT_HEAD, size=12, bold=True, color=BRAND)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        # 管道表格
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-{3,}", c) for c in cells):  # 跳过分隔行
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        cp = t.cell(ri, ci).paragraphs[0]
                        _add_runs(cp, cell)
                        if ri == 0:
                            for r in cp.runs:
                                r.bold = True
            continue
        # 无序/有序列表（含两空格缩进的子项）
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if m:
            indent, marker, content = m.groups()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74 if indent else 0.35)
            p.paragraph_format.space_after = Pt(2)
            bullet = "• " if marker in "-*" else f"{marker} "
            r = p.add_run(bullet)
            _set_font(r, size=10.5)
            _add_runs(p, content)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.4
            _add_runs(p, line)
        i += 1


def _sign(doc: Document, who: str = "学院教务办公室"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"{who}\n{datetime.now().strftime('%Y 年 %m 月 %d 日')}")
    _set_font(r, size=10.5)


def _docx_bytes(doc: Document) -> io.BytesIO:
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# 各类文书
# ---------------------------------------------------------------------------

def letter_docx(letter: dict) -> io.BytesIO:
    """家长信 / 告知函 / 谈话要点：letters.py 的 dict → Word。"""
    doc = _new_doc(letter.get("title", "学业预警文书"))
    _title(doc, letter.get("title", "学业预警文书"), letter.get("subject"))
    md_to_docx(doc, letter.get("body", ""))
    if letter.get("type") == "counselor":
        _sign(doc, who="学生工作办公室")
    else:
        _sign(doc)
    return _docx_bytes(doc)


def notice_docx(engine, sid: str) -> io.BytesIO:
    """预警通知书：比告知函更正式的版式（存档联），含回执条。"""
    s = engine.get(sid)
    r = engine.score(sid)
    doc = _new_doc("学业预警通知书（存档联）")
    _title(doc, "学业预警通知书", f"学号 {s['id']} · {s['class_name']} · {datetime.now().strftime('%Y-%m-%d')}")
    body = [
        f"{s['name']} 同学：",
        "",
        f"经学业预警工作委员会核定，你本学期学业风险综合评分 **{r['score']} 分**，"
        f"达到 **{r['level_label']}** 标准，现予以书面预警。",
        "",
        "核定依据：",
        f"- 课堂缺勤 {r['raw']['absent']} 次 / {r['raw']['sessions']} 次",
        f"- 作业提交 {r['raw']['hw_submitted']} / {r['raw']['hw_assigned']} 次",
        f"- 期中加权均分 {r['raw']['mid_avg']}（班级均分 {r['raw']['cohort_avg']}）",
        f"- 累计不及格 {s['history']['failed_credits']} 学分",
        "",
        "要求：自收到本通知起两周内到辅导员处报到，制定并执行个别化帮扶方案；"
        "学院将在下一教学周复查上述指标。",
        "",
        "（本通知书一式两联，存档联由学院留存，学生联由本人签收。）",
    ]
    md_to_docx(doc, "\n".join(body))
    # 回执条：剪开线
    p = doc.add_paragraph()
    r1 = p.add_run("✂ " + "─" * 34 + " 回执 " + "─" * 34)
    _set_font(r1, size=9, color=MUTED)
    p2 = doc.add_paragraph()
    r2 = p2.add_run(f"本人 {s['name']}（学号 {s['id']}）已收到学业预警通知书，知晓复查要求。\n"
                    f"签收：____________　　日期：____________")
    _set_font(r2, size=10)
    _sign(doc, who="学业预警工作委员会")
    return _docx_bytes(doc)


def interventions_docx(interventions: dict, engine, scope: str = "全部学生") -> io.BytesIO:
    """干预台账：处置留痕，可打印归档。"""
    doc = _new_doc("学业干预台账")
    _title(doc, "学业干预台账", f"{scope} · 截至 {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    rows = [("学生", "班级", "标题", "状态", "备注", "时间")]
    n = 0
    for sid, items in interventions.items():
        try:
            s = engine.get(sid)
            name, cls = s["name"], s["class_name"]
        except KeyError:
            name, cls = sid, "-"
        for it in items:
            rows.append((name, cls, it["title"], it["status"], it.get("note", ""), it["created_at"]))
            n += 1
    t = doc.add_table(rows=len(rows), cols=6)
    t.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            cp = t.cell(ri, ci).paragraphs[0]
            _add_runs(cp, str(cell), base_size=9)
            if ri == 0:
                for r in cp.runs:
                    r.bold = True
    p = doc.add_paragraph()
    r = p.add_run(f"合计 {n} 条干预记录。")
    _set_font(r, size=9, color=MUTED)
    _sign(doc, who="学生工作办公室")
    return _docx_bytes(doc)


def daily_report_docx(report: dict) -> io.BytesIO:
    """每日预警日报：定时任务产出，供早会与推送。"""
    doc = _new_doc("每日学业预警日报")
    _title(doc, "每日学业预警日报",
           f"{report.get('date', '')} · 第 {report.get('week', '-')} 教学周 · 自动生成")
    md_to_docx(doc, report.get("body", ""))
    _sign(doc, who="学业预警智能体（自动生成）")
    return _docx_bytes(doc)


# ---------------------------------------------------------------------------
# 学期汇报 PPT
# ---------------------------------------------------------------------------

def semester_pptx(report: dict, engine) -> io.BytesIO:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn as pqn
    from pptx.util import Inches, Pt as PPt

    BRAND_ = RGBColor(0xC2, 0x41, 0x0C)
    INK_ = RGBColor(0x3B, 0x2E, 0x24)
    MUTED_ = RGBColor(0x8E, 0x7A, 0x66)
    RED_ = RGBColor(0xDC, 0x26, 0x26)

    prs = Presentation()
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height

    def setfont(run, name=FONT_PPT):
        """a:latin 之外补 a:ea（中文字形），PowerPoint/WPS 里中文才不回退。"""
        run.font.name = name
        rPr = run._r.get_or_add_rPr()
        for old in rPr.findall(pqn("a:ea")):
            rPr.remove(old)
        rPr.find(pqn("a:latin")).addnext(
            rPr.makeelement(pqn("a:ea"), {"typeface": name}))

    def add_slide():
        return prs.slides.add_slide(blank)

    def txt(slide, text, x, y, w, h, size=16, color=INK_, bold=False, align=PP_ALIGN.LEFT):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = PPt(size)
        f.color.rgb = color
        f.bold = bold
        setfont(r)
        return tb

    def band(slide, title):
        """每页顶部的标题条。"""
        bar = slide.shapes.add_shape(1, 0, 0, W, Inches(0.9))  # MSO_SHAPE.RECTANGLE
        bar.fill.solid()
        bar.fill.fore_color.rgb = BRAND_
        bar.line.fill.background()
        tf = bar.text_frame
        tf.margin_left = Inches(0.5)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.size = PPt(22)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        setfont(r)

    o = report["overview"]
    dist = o["dist"]

    # 1 封面
    s1 = add_slide()
    band(s1, SCHOOL_NAME)
    txt(s1, "学业预警工作学期汇报", Inches(0.8), Inches(2.3), Inches(8.4), Inches(1.2),
        size=40, bold=True, align=PP_ALIGN.CENTER)
    txt(s1, f"第 {report['week']} 教学周 · {report['date']}", Inches(0.8), Inches(3.6),
        Inches(8.4), Inches(0.5), size=16, color=MUTED_, align=PP_ALIGN.CENTER)
    txt(s1, "数据来源：课堂考勤 / 作业平台 / 期中考核 · 由学业预警智能体自动生成",
        Inches(0.8), Inches(6.3), Inches(8.4), Inches(0.5), size=12, color=MUTED_,
        align=PP_ALIGN.CENTER)

    # 2 全校概览
    s2 = add_slide()
    band(s2, "一、全校风险概览")
    cards = [
        (f"{o['total']}", "在册学生", INK_),
        (f"{dist['red']}", "红色预警", RED_),
        (f"{dist['yellow']}", "黄色关注", RGBColor(0xD9, 0x77, 0x06)),
        (f"{o['avg_score']}", "平均风险分", BRAND_),
    ]
    cw = Inches(2.1)
    for i, (num, label, color) in enumerate(cards):
        x = Inches(0.6) + i * (cw + Inches(0.25))
        txt(s2, num, x, Inches(1.6), cw, Inches(1.1), size=44, color=color, bold=True,
            align=PP_ALIGN.CENTER)
        txt(s2, label, x, Inches(2.75), cw, Inches(0.45), size=15, color=MUTED_,
            align=PP_ALIGN.CENTER)
    txt(s2, "· 红色占比 " + str(report.get("red_pct", "-")) + "%，较规则线（8%）" +
        ("超出，需重点处置" if dist["red"] / max(o["total"], 1) > 0.08 else "在可控区间"),
        Inches(0.6), Inches(3.8), Inches(8.8), Inches(0.5), size=15)
    top_factors = o.get("top_factors", [])[:3]
    if top_factors:
        lines = ["· 全校最普遍的风险成因："] + [
            f"   {f['label']}（均摊 {f['points']:.0f} 分）" for f in top_factors]
        txt(s2, "\n".join(lines), Inches(0.6), Inches(4.5), Inches(8.8), Inches(1.8), size=15)

    # 3 班级排名
    s3 = add_slide()
    band(s3, "二、班级风险排名（红色人数 Top 5）")
    by_class = sorted(o["by_class"].items(), key=lambda kv: -kv[1]["red"])[:5]
    rows = [("班级", "红色", "黄色", "平均分")] + [
        (k, str(v.get("red", 0)), str(v.get("yellow", 0)), str(v.get("avg", "-")))
        for k, v in by_class]
    tbl = s3.shapes.add_table(len(rows), 4, Inches(0.7), Inches(1.4), Inches(8.6),
                              Inches(0.42 * len(rows))).table
    for ci, w in enumerate([Inches(3.4), Inches(1.7), Inches(1.7), Inches(1.8)]):
        tbl.columns[ci].width = w
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            c = tbl.cell(ri, ci)
            c.text = cell
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = PPt(14)
                    setfont(r)
                    r.font.color.rgb = INK_ if ri else BRAND_
                    r.font.bold = ri == 0 or ci == 1 and ri > 0

    # 4 重点学生
    s4 = add_slide()
    band(s4, "三、重点学生（红色预警 Top 6）")
    red_rows = [("学生", "班级", "风险分", "较上学期", "主因")] + [
        (x["name"], x["class_name"], str(x["score"]), f"{x['delta']:+.0f}", x["top_factor"])
        for x in report.get("red_top", [])[:6]]
    tbl4 = s4.shapes.add_table(len(red_rows), 5, Inches(0.7), Inches(1.4), Inches(8.6),
                               Inches(0.42 * len(red_rows))).table
    for ci, w in enumerate([Inches(1.6), Inches(1.8), Inches(1.3), Inches(1.5), Inches(2.4)]):
        tbl4.columns[ci].width = w
    for ri, row in enumerate(red_rows):
        for ci, cell in enumerate(row):
            c = tbl4.cell(ri, ci)
            c.text = cell
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = PPt(13)
                    setfont(r)
                    r.font.color.rgb = INK_ if ri else BRAND_
                    r.font.bold = ri == 0

    # 5 干预与成效
    s5 = add_slide()
    band(s5, "四、干预与成效")
    iv_n = report.get("interventions_total", 0)
    rising_n = len(report.get("rising", []))
    drop_n = report.get("improved_total", 0)
    stats = [
        (f"{iv_n}", "累计干预记录"),
        (f"{rising_n}", "风险上升学生"),
        (f"{drop_n}", "风险回落学生"),
    ]
    for i, (num, label) in enumerate(stats):
        x = Inches(0.9) + i * Inches(3.0)
        txt(s5, num, x, Inches(1.6), Inches(2.6), Inches(1.0), size=40, bold=True,
            color=BRAND_ if i != 1 else RED_, align=PP_ALIGN.CENTER)
        txt(s5, label, x, Inches(2.6), Inches(2.6), Inches(0.45), size=15, color=MUTED_,
            align=PP_ALIGN.CENTER)
    txt(s5, "· 已形成的处置闭环：预警发现 → 对话生成谈话要点 → 处置 → 台账留痕 → 复查",
        Inches(0.6), Inches(3.6), Inches(8.8), Inches(0.5), size=15)

    # 6 下学期建议
    s6 = add_slide()
    band(s6, "五、下阶段工作建议")
    adv = [
        "1. 对红色预警学生完成一轮全覆盖谈话，两周内复查指标；",
        "2. 关注「风险上升但未转红」的学生，提前约谈，避免升级；",
        "3. 按成因分型组织帮扶：缺勤型抓考勤约束，成绩型组织答疑与朋辈辅导；",
        "4. 家校联动：通过家长信与定期反馈，把家庭拉进帮扶链条。",
    ]
    txt(s6, "\n\n".join(adv), Inches(0.8), Inches(1.5), Inches(8.4), Inches(4.5), size=17)
    txt(s6, f"{SCHOOL_NAME} 学业预警智能体 · {report['date']} 自动生成",
        Inches(0.8), Inches(6.6), Inches(8.4), Inches(0.4), size=12, color=MUTED_,
        align=PP_ALIGN.CENTER)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# 邮件：配置了 SMTP 就真发，没配置就产 .eml（双击即进邮件客户端）
# ---------------------------------------------------------------------------

def smtp_configured() -> bool:
    import os
    return bool(os.environ.get("AG_SMTP_HOST"))


def send_letter_email(letter: dict, to_addr: str) -> dict:
    """发送家长信。返回 {mode, ok, ...}；eml 模式返回文件字节。"""
    from email.header import Header
    from email.mime.text import MIMEText
    import os
    import smtplib

    subject = letter.get("subject", letter.get("title", "学业预警"))
    # 正文：markdown 去掉 ** 和 ###，邮件里更干净
    plain = re.sub(r"###?\s*", "", letter.get("body", ""))
    plain = _BOLD_RE.sub(r"\1", plain)
    msg = MIMEText(plain, "plain", "utf-8")
    msg["Subject"] = Header(subject, "utf-8")
    from_name = os.environ.get("AG_SMTP_FROM_NAME", f"{SCHOOL_NAME}学业预警")
    msg["From"] = f"{from_name} <{os.environ.get('AG_SMTP_USER', 'academic-guardian@school.edu.cn')}>"
    msg["To"] = to_addr

    if smtp_configured():
        try:
            host = os.environ["AG_SMTP_HOST"]
            port = int(os.environ.get("AG_SMTP_PORT", "465"))
            user = os.environ.get("AG_SMTP_USER", "")
            pwd = os.environ.get("AG_SMTP_PASS", "")
            with smtplib.SMTP_SSL(host, port, timeout=20) as sm:
                if user:
                    sm.login(user, pwd)
                sm.sendmail(msg["From"], [to_addr], msg.as_string())
            return {"mode": "smtp", "ok": True, "to": to_addr}
        except Exception as exc:  # noqa: BLE001
            return {"mode": "smtp", "ok": False, "error": str(exc)}
    # 演示模式：返回 eml 字节，前端提供下载
    buf = io.BytesIO(msg.as_bytes())
    buf.seek(0)
    return {"mode": "eml", "ok": True, "to": to_addr, "eml": buf}
